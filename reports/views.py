"""reports/views.py — all management reports for ABA Uganda.

Every on-screen report and its PDF download are built from the SAME query
helpers below, so the two can never drift apart (scope, statuses, totals).

Conventions used throughout:
  * Dates use timezone.localdate() (Africa/Kampala), never date.today().
  * Branch scoping is applied in one place (_branch / _effective_branch_id).
  * "Cash received" excludes internal client-credit transfers
    (Payment.objects.cash_receipts()); principal/interest/penalty splits
    include them, because that is when the credit is actually earned.
  * "Overdue" = any unpaid installment past due, i.e. PENDING, OVERDUE *or*
    PARTIAL, measured as total_payment + penalty_due - amount_paid.
"""

import logging
from collections import defaultdict
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
from xml.sax.saxutils import escape

from django.contrib import messages
from django.contrib.auth import get_user_model
from django.contrib.auth.decorators import login_required
from django.db import IntegrityError, transaction
from django.db.models import CharField, Count, DateField, Q, Sum, Case, When
from django.db.models.functions import Abs, Cast
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.utils.dateparse import parse_date
from reportlab.lib.units import mm

from accounts.models import AuditLog, BankTransaction, Branch, CapitalInjection, Expense, SalesAdjustment
from accounts.audit import log_action
from accounts.branch_scope import (
    can_access_branch_object, resolve_branch_for_write, scope_to_branch,
)
from loans.models import Loan, LoanSchedule, LoanProduct
from payments.models import Payment
from clients.models import Client, CreditTransaction
from .pdf_utils import build_report_pdf, p, CELL_BOLD

logger = logging.getLogger("reports")

ZERO = Decimal("0")
OPEN_SCHEDULE = ["PENDING", "OVERDUE", "PARTIAL"]
LIVE_LOAN = ["ACTIVE", "RESTRUCTURED"]
# Loans whose principal has actually been paid out (keep in sync with accounts/cash.py).
DISBURSED = ["ACTIVE", "COMPLETED", "DEFAULTED", "WRITTEN_OFF", "RESTRUCTURED"]


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------
def _ugx(value):
    return f"UGX {value:,.0f}"


def _today():
    return timezone.localdate()


def _stamp():
    return f"{timezone.localtime():%d %b %Y %H:%M}"


def _effective_branch_id(request):
    """Branch a report is scoped to: CEO's ?branch=/switcher, everyone else
    their own branch (-1 = unassigned -> deliberately matches nothing)."""
    from accounts.branch_scope import effective_branch_id
    return effective_branch_id(request)


def _branch(qs, request, field="branch_id"):
    branch_id = _effective_branch_id(request)
    return qs.filter(**{field: branch_id}) if branch_id else qs


def _unpaid(entry):
    """Amount still owed on a schedule row."""
    return max(entry.total_payment + entry.penalty_due - entry.amount_paid, ZERO)


def _sum(qs, field):
    return qs.aggregate(t=Sum(field))["t"] or ZERO


def _period(request):
    today = _today()
    return (
        request.GET.get("date_from", today.replace(day=1).isoformat()),
        request.GET.get("date_to", today.isoformat()),
    )


def _month(request):
    """(month_str, first_day, last_day_inclusive) for the ?month=YYYY-MM picker."""
    today = _today()
    month_str = request.GET.get("month", today.strftime("%Y-%m"))
    try:
        year, month = int(month_str[:4]), int(month_str[5:7])
        first = date(year, month, 1)
    except (ValueError, IndexError):
        first = today.replace(day=1)
        month_str = first.strftime("%Y-%m")
    nxt = date(first.year + (first.month == 12), first.month % 12 + 1, 1)
    return month_str, first, nxt - timedelta(days=1)


def _require_manager(view_fn):
    from functools import wraps

    @wraps(view_fn)
    @login_required
    def wrapper(request, *args, **kwargs):
        if not request.user.can("can_view_reports"):
            messages.error(request, "You do not have permission to view reports.")
            return redirect("accounts:dashboard")
        return view_fn(request, *args, **kwargs)
    return wrapper


# ---------------------------------------------------------------------------
# Shared data builders (used by BOTH the HTML view and the PDF download)
# ---------------------------------------------------------------------------
def _payments_qs(request, date_from, date_to, method=""):
    qs = Payment.objects.filter(
        payment_date__gte=date_from, payment_date__lte=date_to, status="ALLOCATED",
    )
    if method:
        qs = qs.filter(payment_method=method)
    return _branch(qs, request, "loan__branch_id")


def _disbursed_qs(request, date_from, date_to, product_id=""):
    qs = Loan.objects.filter(
        disbursement_date__gte=date_from, disbursement_date__lte=date_to,
        status__in=DISBURSED,
    )
    if product_id:
        qs = qs.filter(product_id=product_id)
    return _branch(qs, request)


def _fees(loans):
    return sum((l.effective_processing_fee for l in loans), ZERO)


def _loan_book_qs(request):
    qs = Loan.objects.filter(status__in=DISBURSED).select_related("client", "product")
    date_from = request.GET.get("date_from", "")
    date_to = request.GET.get("date_to", "")
    product_id = request.GET.get("product", "")
    if date_from:
        qs = qs.filter(disbursement_date__gte=date_from)
    if date_to:
        qs = qs.filter(disbursement_date__lte=date_to)
    if product_id:
        qs = qs.filter(product_id=product_id)
    return _branch(qs, request).order_by("-disbursement_date")


def _collections_data(request):
    date_from, date_to = _period(request)
    payments = (
        _payments_qs(request, date_from, date_to).cash_receipts()
        .select_related("loan__client", "recorded_by")
        .order_by("recorded_by__last_name", "-payment_date")
    )
    by_cashier = defaultdict(list)
    for pm in payments:
        by_cashier[pm.recorded_by].append(pm)
    cashier_totals = [
        {"user": user, "payments": pmts,
         "total": sum((x.amount_received for x in pmts), ZERO), "count": len(pmts)}
        for user, pmts in by_cashier.items()
    ]
    return {
        "date_from": date_from, "date_to": date_to,
        "cashier_totals": cashier_totals,
        "grand_total": sum((r["total"] for r in cashier_totals), ZERO),
        "payment_count": sum(r["count"] for r in cashier_totals),
    }


def _overdue_data(request):
    today = _today()
    entries = LoanSchedule.objects.filter(
        due_date__lt=today, status__in=OPEN_SCHEDULE, loan__status__in=LIVE_LOAN,
    ).select_related("loan__client", "loan__product").order_by("due_date")
    entries = _branch(entries, request, "loan__branch_id")
    rows = []
    for entry in entries:
        amount = _unpaid(entry)
        if amount <= 0:
            continue
        rows.append({
            "entry": entry,
            "days_overdue": (today - entry.due_date).days,
            "overdue_amount": amount,
        })
    rows.sort(key=lambda r: r["days_overdue"], reverse=True)
    return {"rows": rows, "total_overdue": sum((r["overdue_amount"] for r in rows), ZERO), "today": today}


PAR_BUCKETS = [
    # (key, label, css, lower_days, upper_days_exclusive)
    ("par1",  "PAR1 — 1 to 30 Days Overdue",   "par1",  1,  31),
    ("par30", "PAR30 — 31 to 60 Days Overdue", "par30", 31, 61),
    ("par60", "PAR60 — 61 to 90 Days Overdue", "par60", 61, 91),
    ("par90", "PAR90 — Over 90 Days Overdue",  "par90", 91, None),
]


def _par_data(request):
    """Portfolio-at-risk. A loan sits in ONE bucket, chosen by its oldest unpaid
    overdue installment. overdue_amount = unpaid part of overdue installments;
    the % is the loan's full outstanding balance over the active portfolio
    (the standard PAR definition)."""
    today = _today()
    loans = _branch(Loan.objects.filter(status__in=LIVE_LOAN), request)
    per_loan = {}
    entries = LoanSchedule.objects.filter(
        loan__in=loans, due_date__lt=today, status__in=OPEN_SCHEDULE,
    ).order_by("due_date")
    for e in entries:
        rec = per_loan.setdefault(e.loan_id, {"oldest": e.due_date, "amount": ZERO})
        rec["amount"] += _unpaid(e)

    loan_map = {l.pk: l for l in loans.filter(pk__in=list(per_loan)).select_related("client", "product")}
    portfolio = _sum(loans, "outstanding_balance")

    buckets = {key: [] for key, *_ in PAR_BUCKETS}
    for loan_id, rec in per_loan.items():
        if rec["amount"] <= 0 or loan_id not in loan_map:
            continue
        days = (today - rec["oldest"]).days
        for key, _label, _css, lo, hi in PAR_BUCKETS:
            if days >= lo and (hi is None or days < hi):
                buckets[key].append({"loan": loan_map[loan_id], "days_overdue": days,
                                     "overdue_amount": rec["amount"]})
                break

    out = []
    for key, label, css, _lo, _hi in PAR_BUCKETS:
        rows = sorted(buckets[key], key=lambda r: r["days_overdue"], reverse=True)
        at_risk = sum((r["loan"].outstanding_balance for r in rows), ZERO)
        pct = round(float(at_risk) / float(portfolio) * 100, 2) if portfolio else 0
        out.append({"key": key, "label": label, "css": css, "rows": rows,
                    "at_risk": at_risk, "pct": pct})
    return {"buckets": out, "portfolio": portfolio, "today": today}


def _defaulted_data(request):
    today = _today()
    qs = _branch(
        Loan.objects.filter(status__in=["DEFAULTED", "WRITTEN_OFF"])
        .select_related("client", "product", "reviewed_by").order_by("-updated_at"),
        request,
    )
    loans = list(qs)

    # Write-off zeroes outstanding_balance, so recover the amount from the
    # installments the write-off waived.
    wo_ids = [l.pk for l in loans if l.status == "WRITTEN_OFF"]
    wo_amounts = defaultdict(lambda: ZERO)
    for e in LoanSchedule.objects.filter(loan_id__in=wo_ids, waived_by_writeoff=True):
        wo_amounts[e.loan_id] += _unpaid(e)
    for l in loans:
        l.written_off_amount = wo_amounts.get(l.pk, ZERO) if l.status == "WRITTEN_OFF" else ZERO

    par90 = [r["loan"] for r in next(b for b in _par_data(request)["buckets"] if b["key"] == "par90")["rows"]]
    return {
        "loans": loans,
        "par90_loans": par90,
        "total_defaulted": sum((l.principal_amount for l in loans), ZERO),
        "total_outstanding": sum((l.outstanding_balance for l in loans), ZERO),
        "total_written_off": sum((l.written_off_amount for l in loans), ZERO),
        "loan_count": len(loans),
        "par90_count": len(par90),
        "today": today,
    }


def _income_data(request):
    month_str, first, last = _month(request)
    payments = _payments_qs(request, first, last)
    cash_payments = payments.cash_receipts()

    total_received = _sum(cash_payments, "amount_received")
    total_principal = _sum(payments, "principal_paid")
    total_interest = _sum(payments, "interest_paid")
    total_penalties = _sum(payments, "penalty_paid")
    total_allocated = total_principal + total_interest + total_penalties

    disbursed = list(_disbursed_qs(request, first, last).select_related("product"))
    total_fees = _fees(disbursed)
    expenses = _branch(Expense.objects.filter(
        expense_date__gte=first, expense_date__lte=last, status="APPROVED"), request)
    total_expenses = _sum(expenses, "amount")

    total_income = total_interest + total_penalties + total_fees
    return {
        "month_str": month_str, "period_start": first, "period_end": last,
        "total_received": total_received,
        "total_principal": total_principal,
        "total_interest": total_interest,
        "total_penalties": total_penalties,
        "total_fees": total_fees,
        "total_income": total_income,
        "total_expenses": total_expenses,
        "net_income": total_income - total_expenses,
        # >0: client credit was consumed this month; <0: overpayments were held as credit.
        "credit_adjustment": total_allocated - total_received,
        "total_disbursed": sum((l.principal_amount for l in disbursed), ZERO),
        "loan_count": len(disbursed),
        "payment_count": cash_payments.count(),
    }


def _cash_flow_data(request):
    date_from, date_to = _period(request)
    sort_by = request.GET.get("sort", "date")
    if sort_by not in {"date", "cash_in", "cash_out"}:
        sort_by = "date"

    payments = list(_branch(
        Payment.objects.cash_receipts().filter(
            payment_date__gte=date_from, payment_date__lte=date_to, status="ALLOCATED",
            payment_method=Payment.PaymentMethod.CASH),
        request, "loan__branch_id",
    ).select_related("loan__client", "client", "recorded_by").order_by("payment_date"))

    disbursements = list(_disbursed_qs(request, date_from, date_to)
                         .select_related("client", "product").order_by("disbursement_date"))

    expenses = list(_branch(Expense.objects.filter(
        expense_date__gte=date_from, expense_date__lte=date_to, status="APPROVED",
        payment_method=Expense.PaymentMethod.CASH,
    ), request).select_related("category", "expense_type").order_by("expense_date"))

    injections = list(_branch(CapitalInjection.objects.filter(
        injected_date__gte=date_from, injected_date__lte=date_to,
        payment_method=CapitalInjection.PaymentMethod.CASH,
    ), request).order_by("injected_date"))

    total_processing_fees = _fees(disbursements)
    total_expenses = sum((e.amount for e in expenses), ZERO)
    total_cash_in = (sum((x.amount_received for x in payments), ZERO)
                     + sum((i.amount for i in injections), ZERO) + total_processing_fees)
    total_cash_out = sum((l.cash_disbursed for l in disbursements), ZERO) + total_expenses

    def row(**kw):
        base = {"loan": None, "client": None, "cash_in": ZERO, "cash_out": ZERO,
                "processing_fee": ZERO, "principal": ZERO, "interest": ZERO, "penalty": ZERO}
        base.update(kw)
        return base

    ledger = []
    for pm in payments:
        ledger.append(row(
            date=pm.payment_date, time=pm.created_at, type="Payment", loan=pm.loan, client=pm.client,
            cash_in=pm.amount_received, principal=pm.principal_paid, interest=pm.interest_paid,
            penalty=pm.penalty_paid, sort_order=1,
            description=f"Payment received ({pm.get_payment_method_display()})"))
    for seq, l in enumerate(disbursements):
        if l.cash_disbursed:
            ledger.append(row(
                date=l.disbursement_date, time=l.created_at, type="Disbursement", loan=l, client=l.client,
                cash_out=l.cash_disbursed, principal=l.principal_amount, sort_order=3 + seq * 2,
                description="Principal disbursed to client"))
        fee = l.effective_processing_fee
        if fee:
            ledger.append(row(
                date=l.disbursement_date, time=l.created_at, type="Processing Fee", loan=l, client=l.client,
                cash_in=fee, sort_order=4 + seq * 2, description="Processing fee collected"))
    for e in expenses:
        category = e.expense_type.name if e.expense_type else (e.category.name if e.category else "Expense")
        ledger.append(row(
            date=e.expense_date, time=e.created_at, type="Expense", cash_out=e.amount, sort_order=2,
            description=category + (f" — {e.vendor}" if e.vendor else "")))
    for i in injections:
        ledger.append(row(
            date=i.injected_date, time=i.created_at, type="Capital Injection", cash_in=i.amount,
            sort_order=0, description=f"Capital injection from {i.source}"))

    if sort_by == "cash_in":
        ledger.sort(key=lambda r: (-r["cash_in"], r["date"], r["sort_order"]))
    elif sort_by == "cash_out":
        ledger.sort(key=lambda r: (-r["cash_out"], r["date"], r["sort_order"]))
    else:
        ledger.sort(key=lambda r: (r["date"], r["sort_order"], r["time"], r["type"]))

    return {
        "date_from": date_from, "date_to": date_to, "ledger": ledger, "sort_by": sort_by,
        "total_cash_in": total_cash_in, "total_cash_out": total_cash_out,
        "total_processing_fees": total_processing_fees, "total_expenses": total_expenses,
        "total_principal": sum((x.principal_paid for x in payments), ZERO),
        "total_interest": sum((x.interest_paid for x in payments), ZERO),
        "total_penalties": sum((x.penalty_paid for x in payments), ZERO),
        "payment_count": len(payments), "disbursement_count": len(disbursements),
        "expense_count": len(expenses),
    }


def _fund_movement_data(request):
    """Build a reconciled, all-channel fund ledger for the selected period."""
    today = _today()
    injection_scope = _branch(CapitalInjection.objects.all(), request)
    earliest_injection = injection_scope.order_by("injected_date").values_list(
        "injected_date", flat=True
    ).first()
    default_from = earliest_injection or today.replace(day=1)
    date_from = parse_date(request.GET.get("date_from", "")) or default_from
    date_to = parse_date(request.GET.get("date_to", "")) or today
    date_to = min(date_to, today)
    if date_from > date_to:
        date_from = date_to
    branch_id = _effective_branch_id(request)
    branch_label = (
        Branch.objects.filter(pk=branch_id).values_list("name", flat=True).first() or "Selected branch"
        if branch_id else "All branches"
    )

    payments = list(_branch(
        Payment.objects.cash_receipts().filter(
            payment_date__gte=date_from,
            payment_date__lte=date_to,
            status__in=[Payment.Status.ALLOCATED, Payment.Status.REVERSED],
        ), request, "loan__branch_id",
    ).select_related("loan__branch", "client").order_by("payment_date", "created_at"))
    payment_reversals = list(_branch(
        Payment.objects.cash_receipts().filter(
            status=Payment.Status.REVERSED,
        ).filter(
            Q(reversal_date__gte=date_from, reversal_date__lte=date_to)
            | Q(
                reversal_date__isnull=True,
                payment_date__gte=date_from,
                payment_date__lte=date_to,
            ),
        ), request, "loan__branch_id",
    ).select_related("loan__branch", "client", "reversed_by").order_by("reversal_date", "created_at"))
    refunds = list(_branch(
        CreditTransaction.objects.filter(
            tx_type=CreditTransaction.TxType.REFUNDED,
            created_at__date__gte=date_from,
            created_at__date__lte=date_to,
        ), request, "client__branch_id",
    ).select_related("client", "client__branch", "created_by").order_by("created_at"))
    injections = list(_branch(
        CapitalInjection.objects.filter(injected_date__gte=date_from, injected_date__lte=date_to),
        request,
    ).select_related("branch", "created_by").order_by("injected_date", "created_at"))
    expenses = list(_branch(
        Expense.objects.filter(
            expense_date__gte=date_from,
            expense_date__lte=date_to,
            status=Expense.Status.APPROVED,
        ), request,
    ).select_related("branch", "category", "expense_type", "approved_by").order_by("expense_date", "created_at"))
    adjustments = list(_branch(
        SalesAdjustment.objects.filter(
            adjustment_date__gte=date_from,
            adjustment_date__lte=date_to,
        ), request,
    ).select_related("branch", "created_by").order_by("adjustment_date", "created_at"))
    disbursements = list(_disbursed_qs(request, date_from, date_to)
                         .select_related("client", "product", "branch")
                         .order_by("disbursement_date", "created_at"))

    prior_payments = _branch(Payment.objects.cash_receipts().filter(
        payment_date__lt=date_from,
        status__in=[Payment.Status.ALLOCATED, Payment.Status.REVERSED],
    ).filter(
        Q(status=Payment.Status.ALLOCATED)
        | Q(reversal_date__gte=date_from),
    ), request, "loan__branch_id")
    prior_injections = _branch(CapitalInjection.objects.filter(injected_date__lt=date_from), request)
    prior_expenses = _branch(Expense.objects.filter(
        expense_date__lt=date_from, status=Expense.Status.APPROVED,
    ), request)
    prior_adjustments = _branch(SalesAdjustment.objects.filter(
        adjustment_date__lt=date_from,
    ), request)
    prior_refunds = _branch(CreditTransaction.objects.filter(
        tx_type=CreditTransaction.TxType.REFUNDED,
        created_at__date__lt=date_from,
    ), request, "client__branch_id")
    prior_disbursements = list(_disbursed_qs(request, date_from="1900-01-01", date_to=date_from - timedelta(days=1))
                               .select_related("product"))

    opening_balance = (
        _sum(prior_payments, "amount_received")
        + _sum(prior_injections, "amount")
        + sum((loan.effective_processing_fee for loan in prior_disbursements), ZERO)
        + _sum(prior_adjustments.filter(direction=SalesAdjustment.Direction.CASH_IN), "amount")
        - _sum(prior_expenses, "amount")
        - (prior_refunds.aggregate(total=Sum(Abs("amount")))["total"] or ZERO)
        - _sum(prior_adjustments.filter(direction=SalesAdjustment.Direction.CASH_OUT), "amount")
        - sum((loan.principal_amount for loan in prior_disbursements), ZERO)
    )

    ledger = []

    def add_row(*, event_date, event_time, kind, counterparty, method, reference,
                description, cash_in=ZERO, cash_out=ZERO, branch="", sort_order=0):
        ledger.append({
            "date": event_date,
            "time": event_time,
            "type": kind,
            "counterparty": counterparty or "—",
            "method": method or "—",
            "reference": reference or "—",
            "description": description or "",
            "cash_in": cash_in,
            "cash_out": cash_out,
            "branch": branch or "—",
            "sort_order": sort_order,
        })

    for payment in payments:
        add_row(
            event_date=payment.payment_date,
            event_time=payment.created_at,
            kind="Repayment",
            counterparty=payment.client.full_name,
            method=payment.get_payment_method_display(),
            reference=payment.reference_number or payment.loan.loan_number,
            description=(
                f"Loan {payment.loan.loan_number}: principal {_ugx(payment.principal_paid)}, "
                f"interest {_ugx(payment.interest_paid)}, penalty {_ugx(payment.penalty_paid)}"
            ),
            cash_in=payment.amount_received,
            branch=payment.loan.branch.name if payment.loan.branch else "",
            sort_order=1,
        )
    for payment in payment_reversals:
        add_row(
            event_date=payment.reversal_date or payment.payment_date,
            event_time=payment.updated_at,
            kind="Payment reversal",
            counterparty=payment.client.full_name,
            method=payment.get_payment_method_display(),
            reference=payment.reference_number or payment.loan.loan_number,
            description=payment.reversal_reason or f"Reversed payment for {payment.loan.loan_number}",
            cash_out=payment.amount_received,
            branch=payment.loan.branch.name if payment.loan.branch else "",
            sort_order=2,
        )
    for refund in refunds:
        add_row(
            event_date=timezone.localtime(refund.created_at).date(),
            event_time=refund.created_at,
            kind="Client refund",
            counterparty=refund.client.full_name,
            method="Client credit",
            reference=f"CR-{refund.pk}",
            description=refund.notes,
            cash_out=abs(refund.amount),
            branch=refund.client.branch.name if refund.client.branch else "",
            sort_order=2,
        )
    for injection in injections:
        add_row(
            event_date=injection.injected_date,
            event_time=injection.created_at,
            kind="Capital injection",
            counterparty=injection.investor or injection.source,
            method=injection.get_payment_method_display(),
            reference=f"INJ-{injection.pk}",
            description=injection.notes or injection.source,
            cash_in=injection.amount,
            branch=injection.branch.name if injection.branch else "",
            sort_order=0,
        )
    for expense in expenses:
        category = expense.expense_type.name if expense.expense_type else (
            expense.category.name if expense.category else "Operating expense"
        )
        add_row(
            event_date=expense.expense_date,
            event_time=expense.created_at,
            kind="Expense",
            counterparty=expense.vendor or "—",
            method=expense.get_payment_method_display(),
            reference=expense.receipt_number or expense.reference_number,
            description=f"{category}: {expense.description}".strip(": "),
            cash_out=expense.amount,
            branch=expense.branch.name if expense.branch else "",
            sort_order=2,
        )
    for adjustment in adjustments:
        add_row(
            event_date=adjustment.adjustment_date,
            event_time=adjustment.created_at,
            kind="Sales adjustment",
            counterparty=adjustment.created_by.get_full_name() if adjustment.created_by else "—",
            method=adjustment.get_direction_display(),
            reference=adjustment.reference,
            description=adjustment.notes,
            cash_in=adjustment.amount if adjustment.direction == SalesAdjustment.Direction.CASH_IN else ZERO,
            cash_out=adjustment.amount if adjustment.direction == SalesAdjustment.Direction.CASH_OUT else ZERO,
            branch=adjustment.branch.name if adjustment.branch else "",
            sort_order=2,
        )
    for loan in disbursements:
        add_row(
            event_date=loan.disbursement_date,
            event_time=loan.created_at,
            kind="Loan disbursement",
            counterparty=loan.client.full_name,
            method="Not recorded",
            reference=loan.loan_number,
            description=f"{loan.product.name} principal disbursed",
            cash_out=loan.principal_amount,
            branch=loan.branch.name if loan.branch else "",
            sort_order=3,
        )
        fee = loan.effective_processing_fee
        if fee:
            add_row(
                event_date=loan.disbursement_date,
                event_time=loan.created_at,
                kind="Processing fee",
                counterparty=loan.client.full_name,
                method="Collected at disbursement",
                reference=loan.loan_number,
                description=f"Processing fee for {loan.product.name}",
                cash_in=fee,
                branch=loan.branch.name if loan.branch else "",
                sort_order=4,
            )

    ledger.sort(key=lambda row: (row["date"], row["time"], row["sort_order"], row["type"]))
    running_balance = opening_balance
    for row in ledger:
        running_balance += row["cash_in"] - row["cash_out"]
        row["balance_after"] = running_balance

    total_injections = sum((item.amount for item in injections), ZERO)
    total_repayments = sum((item.amount_received for item in payments), ZERO)
    total_processing_fees = sum((loan.effective_processing_fee for loan in disbursements), ZERO)
    total_disbursed = sum((loan.principal_amount for loan in disbursements), ZERO)
    total_expenses = sum((item.amount for item in expenses), ZERO)
    total_adjustment_in = sum(
        (item.amount for item in adjustments if item.direction == SalesAdjustment.Direction.CASH_IN), ZERO,
    )
    total_adjustment_out = sum(
        (item.amount for item in adjustments if item.direction == SalesAdjustment.Direction.CASH_OUT), ZERO,
    )
    total_refunds = sum((abs(item.amount) for item in refunds), ZERO)
    total_reversals = sum((item.amount_received for item in payment_reversals), ZERO)
    total_inflows = total_injections + total_repayments + total_processing_fees + total_adjustment_in
    total_outflows = total_disbursed + total_expenses + total_adjustment_out + total_refunds + total_reversals

    return {
        "date_from": date_from,
        "date_to": date_to,
        "branch_label": branch_label,
        "opening_balance": opening_balance,
        "total_injections": total_injections,
        "total_repayments": total_repayments,
        "total_processing_fees": total_processing_fees,
        "total_adjustment_in": total_adjustment_in,
        "total_adjustment_out": total_adjustment_out,
        "total_refunds": total_refunds,
        "total_reversals": total_reversals,
        "total_inflows": total_inflows,
        "total_disbursed": total_disbursed,
        "total_expenses": total_expenses,
        "total_outflows": total_outflows,
        "net_movement": total_inflows - total_outflows,
        "closing_balance": opening_balance + total_inflows - total_outflows,
        "payment_count": len(payments),
        "injection_count": len(injections),
        "expense_count": len(expenses),
        "adjustment_count": len(adjustments),
        "refund_count": len(refunds),
        "reversal_count": len(payment_reversals),
        "disbursement_count": len(disbursements),
        "ledger": ledger,
    }


def _staff_rows(request, date_from, date_to):
    today = _today()
    branch_id = _effective_branch_id(request)
    role = request.GET.get("role", "")
    staff_id = request.GET.get("staff", "")

    User = get_user_model()
    staff_users = User.objects.filter(is_active=True).order_by("last_name", "first_name")
    if branch_id:
        staff_users = staff_users.filter(branch_id=branch_id)
    if role:
        staff_users = staff_users.filter(role=role)
    if staff_id:
        staff_users = staff_users.filter(pk=staff_id)

    def scoped(qs):
        return qs.filter(branch_id=branch_id) if branch_id else qs

    rows = []
    for user in staff_users:
        disbursed = scoped(Loan.objects.filter(
            disbursement_date__gte=date_from, disbursement_date__lte=date_to,
            applied_by=user, status__in=DISBURSED))
        collected = Payment.objects.cash_receipts().filter(
            payment_date__gte=date_from, payment_date__lte=date_to,
            recorded_by=user, status="ALLOCATED")
        if branch_id:
            collected = collected.filter(loan__branch_id=branch_id)
        overdue = scoped(Loan.objects.filter(
            applied_by=user, status__in=LIVE_LOAN,
            schedule__due_date__lt=today, schedule__status__in=OPEN_SCHEDULE,
        )).distinct()
        overdue_count = overdue.count()
        overdue_amount = sum((l.outstanding_balance for l in overdue), ZERO)
        total_loans = scoped(Loan.objects.filter(applied_by=user)).count()
        defaulted = scoped(Loan.objects.filter(
            applied_by=user, status__in=["DEFAULTED", "WRITTEN_OFF"])).count()
        disbursed_count = disbursed.count()

        quality = Decimal("100")
        if disbursed_count:
            quality -= 5 * min(defaulted, 10)
        if overdue_count:
            quality -= 2 * min(overdue_count, 10)
        quality = max(ZERO, quality)

        rows.append({
            "user": user,
            "disbursed_amount": _sum(disbursed, "principal_amount"),
            "disbursed_count": disbursed_count,
            "collected_amount": _sum(collected, "amount_received"),
            "collected_count": collected.count(),
            "overdue_count": overdue_count, "overdue_amount": overdue_amount,
            "total_loans": total_loans,
            "defaulted_loans": defaulted, "defaulted_count": defaulted,
            "quality_score": quality, "performance_percentage": quality,
        })
    rows.sort(key=lambda r: (-r["disbursed_amount"], r["user"].last_name))
    return rows


# ---------------------------------------------------------------------------
# Index
# ---------------------------------------------------------------------------
LEDGER_CATEGORIES = {
    "payment": "Payments and reversals",
    "refund": "Client refunds",
    "credit": "Client credit movements",
    "injection": "Capital injections",
    "expense": "Expenses",
    "loan": "Loan disbursements",
    "client": "Client registrations",
    "product": "Loan products",
    "bank": "Bank transactions",
    "adjustment": "Sales adjustments",
    "audit": "Accountability activity",
}
LEDGER_PAGE_SIZE = 100


@_require_manager
def company_ledger(request):
    today = _today()
    errors = []
    form_values = {
        "reference": request.POST.get("reference", ""),
        "direction": request.POST.get("direction", SalesAdjustment.Direction.CASH_OUT),
        "amount": request.POST.get("amount", ""),
        "adjustment_date": request.POST.get("adjustment_date", today.isoformat()),
        "adjustment_branch": request.POST.get("adjustment_branch", ""),
        "notes": request.POST.get("notes", ""),
    }
    if request.method == "POST":
        if not request.user.can("can_manage_expenses"):
            return HttpResponse("You do not have permission to record sales adjustments.", status=403)
        reference = form_values["reference"].strip()
        if not reference:
            errors.append("A reference is required for every sales adjustment.")
        elif SalesAdjustment.objects.filter(reference=reference).exists():
            errors.append("That reference is already in use.")
        if form_values["direction"] not in SalesAdjustment.Direction.values:
            errors.append("Choose whether the adjustment is cash in or cash out.")
        try:
            amount = Decimal(form_values["amount"])
            if not amount.is_finite() or amount <= ZERO:
                raise InvalidOperation
        except (InvalidOperation, TypeError, ValueError):
            amount = ZERO
            errors.append("Enter an adjustment amount greater than zero.")
        adjustment_date = parse_date(form_values["adjustment_date"])
        if adjustment_date is None or adjustment_date > today:
            errors.append("Enter a valid adjustment date that is today or earlier.")
        try:
            write_branch_id = resolve_branch_for_write(request, form_values["adjustment_branch"])
            if write_branch_id is None or not Branch.objects.filter(
                pk=write_branch_id, is_active=True,
            ).exists():
                errors.append("Choose an active branch for this adjustment.")
        except ValueError as error:
            errors.append(str(error))
            write_branch_id = None
        if not errors:
            try:
                with transaction.atomic():
                    adjustment = SalesAdjustment.objects.create(
                        reference=reference,
                        direction=form_values["direction"],
                        amount=amount,
                        adjustment_date=adjustment_date,
                        branch_id=write_branch_id,
                        notes=form_values["notes"].strip(),
                        created_by=request.user,
                    )
                    log_action(
                        request.user,
                        AuditLog.Action.CREATE,
                        adjustment,
                        changes={
                            "reference": reference,
                            "direction": adjustment.get_direction_display(),
                            "amount": str(amount),
                        },
                        request=request,
                        remarks=f"Sales adjustment {reference} recorded",
                    )
            except IntegrityError:
                if SalesAdjustment.objects.filter(reference=reference).exists():
                    errors.append("That reference is already in use.")
                else:
                    raise
            else:
                messages.success(request, "Sales adjustment recorded.")
                target = reverse("reports:company_ledger")
                if request.GET:
                    target = f"{target}?{request.GET.urlencode()}"
                return redirect(target)

    default_from = today.replace(day=1)
    raw_from = request.GET.get("date_from") or default_from.isoformat()
    raw_to = request.GET.get("date_to") or today.isoformat()
    date_from = parse_date(raw_from)
    date_to = parse_date(raw_to)
    if date_from is None or date_to is None:
        errors.append("Enter valid start and end dates.")
        date_from, date_to = default_from, today
    elif date_from > date_to:
        errors.append("The start date must be on or before the end date.")
        date_from, date_to = default_from, today
    if date_to > today:
        errors.append("The end date cannot be in the future.")
        date_to = today
        if date_from > date_to:
            date_from = default_from

    branch_id = _effective_branch_id(request)
    raw_category = request.GET.get("category", "")
    visible_categories = dict(LEDGER_CATEGORIES)
    if not request.user.can("can_view_audit_log"):
        visible_categories.pop("audit")
    if not (
        request.user.can("can_manage_bank_accounts")
        or request.user.can("can_view_financial_overview")
    ):
        visible_categories.pop("bank")
    category = raw_category if raw_category in visible_categories else ""
    if raw_category and not category:
        errors.append("Choose a valid ledger category.")
    search = request.GET.get("q", "").strip()[:100]

    def search_in(queryset, *fields):
        if not search:
            return queryset
        condition = Q()
        for field in fields:
            condition |= Q(**{f"{field}__icontains": search})
        return queryset.filter(condition)

    payment_period = (
        Q(payment_date__gte=date_from, payment_date__lte=date_to)
        & ~Q(status=Payment.Status.REVERSED)
    ) | Q(
        status=Payment.Status.REVERSED,
        reversal_date__gte=date_from,
        reversal_date__lte=date_to,
    ) | Q(
        status=Payment.Status.REVERSED,
        reversal_date__isnull=True,
        payment_date__gte=date_from,
        payment_date__lte=date_to,
    )
    payments = _branch(Payment.objects.filter(payment_period).annotate(
        ledger_date=Case(
            When(
                status=Payment.Status.REVERSED,
                reversal_date__isnull=False,
                then="reversal_date",
            ),
            default="payment_date",
            output_field=DateField(),
        ),
    ), request, "loan__branch_id").select_related(
        "loan__branch", "loan", "client", "recorded_by", "reversed_by",
    )
    payments = search_in(
        payments, "reference_number", "loan__loan_number", "client__first_name",
        "client__last_name", "recorded_by__first_name", "recorded_by__last_name",
    )

    credits = _branch(CreditTransaction.objects.filter(
        created_at__date__gte=date_from, created_at__date__lte=date_to,
    ), request, "client__branch_id").select_related(
        "client", "client__branch", "loan", "created_by",
    )
    credits = search_in(
        credits, "client__client_number", "client__first_name", "client__last_name",
        "notes", "tx_type",
    )

    injections = _branch(CapitalInjection.objects.filter(
        injected_date__gte=date_from, injected_date__lte=date_to,
    ), request).select_related("branch", "created_by")
    injections = search_in(injections, "source", "investor", "notes", "created_by__first_name",
                           "created_by__last_name")

    expenses = _branch(Expense.objects.filter(
        expense_date__gte=date_from, expense_date__lte=date_to,
    ), request).select_related("branch", "category", "expense_type", "created_by", "approved_by")
    expenses = search_in(
        expenses, "reference_number", "receipt_number", "vendor", "description",
        "created_by__first_name", "created_by__last_name",
    )

    loans = _branch(Loan.objects.filter(
        disbursement_date__gte=date_from, disbursement_date__lte=date_to,
        status__in=DISBURSED,
    ), request).select_related("branch", "client", "product", "applied_by")
    loans = search_in(
        loans, "loan_number", "client__first_name", "client__last_name",
        "product__name", "applied_by__first_name", "applied_by__last_name",
    )

    clients = _branch(Client.objects.filter(
        created_at__date__gte=date_from, created_at__date__lte=date_to,
    ), request).select_related("branch", "registered_by")
    clients = search_in(
        clients, "client_number", "first_name", "last_name", "other_names",
        "registered_by__first_name", "registered_by__last_name",
    )

    products = LoanProduct.objects.filter(
        created_at__date__gte=date_from, created_at__date__lte=date_to,
    ).select_related("created_by")
    products = search_in(products, "name", "description", "created_by__first_name",
                         "created_by__last_name")

    bank_transactions = BankTransaction.objects.none()
    if request.user.can("can_manage_bank_accounts") or request.user.can("can_view_financial_overview"):
        bank_transactions = _branch(BankTransaction.objects.filter(
            transaction_date__date__gte=date_from, transaction_date__date__lte=date_to,
        ), request).select_related("branch", "bank_account", "created_by")
    bank_transactions = search_in(
        bank_transactions, "reference_number", "category", "notes",
        "bank_account__account_name", "created_by__first_name", "created_by__last_name",
    )

    adjustments = _branch(SalesAdjustment.objects.filter(
        adjustment_date__gte=date_from, adjustment_date__lte=date_to,
    ), request).select_related("branch", "created_by")
    adjustments = search_in(
        adjustments, "reference", "notes", "created_by__first_name", "created_by__last_name",
    )

    audit_logs = AuditLog.objects.none()
    if request.user.can("can_view_audit_log"):
        audit_logs = AuditLog.objects.filter(
            created_at__date__gte=date_from, created_at__date__lte=date_to,
        ).select_related("user")
        if branch_id is not None:
            scoped_entities = (
                ("Loan", Loan.objects.filter(branch_id=branch_id).annotate(
                    entity_pk=Cast("pk", output_field=CharField()),
                ).values("entity_pk")),
                ("Client", Client.objects.filter(branch_id=branch_id).annotate(
                    entity_pk=Cast("pk", output_field=CharField()),
                ).values("entity_pk")),
                ("Payment", Payment.objects.filter(loan__branch_id=branch_id).annotate(
                    entity_pk=Cast("pk", output_field=CharField()),
                ).values("entity_pk")),
                ("Expense", Expense.objects.filter(branch_id=branch_id).annotate(
                    entity_pk=Cast("pk", output_field=CharField()),
                ).values("entity_pk")),
                ("CapitalInjection", CapitalInjection.objects.filter(branch_id=branch_id).annotate(
                    entity_pk=Cast("pk", output_field=CharField()),
                ).values("entity_pk")),
                ("BankTransaction", BankTransaction.objects.filter(branch_id=branch_id).annotate(
                    entity_pk=Cast("pk", output_field=CharField()),
                ).values("entity_pk")),
                ("SalesAdjustment", SalesAdjustment.objects.filter(branch_id=branch_id).annotate(
                    entity_pk=Cast("pk", output_field=CharField()),
                ).values("entity_pk")),
            )
            scope_query = Q(pk__in=[])
            for entity_type, entity_ids in scoped_entities:
                scope_query |= Q(entity_type=entity_type, entity_id__in=entity_ids)
            audit_logs = audit_logs.filter(scope_query)
        audit_logs = search_in(audit_logs, "entity_type", "entity_repr", "action", "remarks")

    querysets = {
        "payment": (payments, "ledger_date"),
        "refund": (credits.filter(tx_type=CreditTransaction.TxType.REFUNDED), "created_at"),
        "credit": (credits.exclude(tx_type=CreditTransaction.TxType.REFUNDED), "created_at"),
        "injection": (injections, "injected_date"),
        "expense": (expenses, "expense_date"),
        "loan": (loans, "disbursement_date"),
        "client": (clients, "created_at"),
        "product": (products, "created_at"),
        "bank": (bank_transactions, "transaction_date"),
        "adjustment": (adjustments, "adjustment_date"),
        "audit": (audit_logs, "created_at"),
    }
    selected_querysets = (
        {category: querysets[category]} if category else querysets
    )

    try:
        page_number = max(1, int(request.GET.get("page", "1")))
    except ValueError:
        page_number = 1
    start = (page_number - 1) * LEDGER_PAGE_SIZE
    end = start + LEDGER_PAGE_SIZE
    total_count = sum(queryset.count() for queryset, _date_field in selected_querysets.values())
    page_count = max(1, (total_count + LEDGER_PAGE_SIZE - 1) // LEDGER_PAGE_SIZE)
    if page_number > page_count:
        page_number = page_count
        start = (page_number - 1) * LEDGER_PAGE_SIZE
        end = start + LEDGER_PAGE_SIZE

    rows = []

    def add_row(kind, event_date, timestamp, reference, description, branch, user,
                amount=ZERO, cash_in=ZERO, cash_out=ZERO, status=""):
        rows.append({
            "category": kind,
            "type": LEDGER_CATEGORIES[kind],
            "date": event_date,
            "time": timestamp,
            "reference": reference or "—",
            "description": description or "",
            "branch": branch or "—",
            "user": user.get_full_name() if user else "—",
            "amount": amount or ZERO,
            "cash_in": cash_in or ZERO,
            "cash_out": cash_out or ZERO,
            "status": status,
        })

    for kind, (queryset, date_field) in selected_querysets.items():
        for item in queryset.order_by(f"-{date_field}", "-created_at")[:end]:
            event_time = getattr(item, date_field)
            event_date = (
                timezone.localtime(event_time).date()
                if hasattr(event_time, "tzinfo") and event_time.tzinfo
                else event_time.date() if hasattr(event_time, "date") else event_time
            )
            branch = None
            user = None
            amount = ZERO
            cash_in = cash_out = ZERO
            reference = description = status = ""
            if kind == "payment":
                branch = item.loan.branch.name if item.loan.branch else None
                user = item.reversed_by if item.status == Payment.Status.REVERSED else item.recorded_by
                amount = item.amount_received
                cash_in = amount if item.status == Payment.Status.ALLOCATED else ZERO
                cash_out = amount if item.status == Payment.Status.REVERSED else ZERO
                reference = item.reference_number or item.loan.loan_number
                description = f"{item.client.full_name} · Loan {item.loan.loan_number}"
                status = item.get_status_display()
            elif kind in {"refund", "credit"}:
                branch = item.client.branch.name if item.client.branch else None
                user = item.created_by
                amount = abs(item.amount)
                cash_out = amount if item.tx_type == CreditTransaction.TxType.REFUNDED else ZERO
                reference = f"CR-{item.pk}"
                description = f"{item.client.full_name} · {item.get_tx_type_display()}: {item.notes}".strip()
                status = item.get_tx_type_display()
            elif kind == "injection":
                branch, user, amount = item.branch, item.created_by, item.amount
                cash_in = amount
                reference = f"INJ-{item.pk}"
                description = f"{item.investor or item.source}: {item.notes}".strip(": ")
                status = item.get_payment_method_display()
            elif kind == "expense":
                branch, user, amount = item.branch, item.created_by, item.amount
                cash_out = amount if item.status == Expense.Status.APPROVED else ZERO
                reference = item.receipt_number or item.reference_number
                description = f"{item.vendor or 'Expense'} · {item.description}".strip(" ·")
                status = item.get_status_display()
            elif kind == "loan":
                branch, user, amount = item.branch, item.applied_by, item.principal_amount
                cash_out = amount
                reference = item.loan_number
                description = f"{item.client.full_name} · {item.product.name}"
                status = item.get_status_display()
            elif kind == "client":
                branch, user = item.branch, item.registered_by
                reference = item.client_number
                description = item.full_name
                status = "Active" if item.is_active else "Inactive"
            elif kind == "product":
                user = item.created_by
                reference = f"PRODUCT-{item.pk}"
                description = item.name
                status = "Active" if item.is_active else "Inactive"
            elif kind == "bank":
                branch, user, amount = item.branch, item.created_by, item.amount
                reference = item.reference_number or f"BANK-{item.pk}"
                description = f"{item.bank_account.account_name} · {item.category or item.notes}"
                status = item.get_transaction_type_display()
                if item.transaction_type == BankTransaction.TransactionType.CREDIT:
                    cash_in = amount
                else:
                    cash_out = amount
            elif kind == "adjustment":
                branch, user, amount = item.branch, item.created_by, item.amount
                reference = item.reference
                description = item.notes
                status = item.get_direction_display()
                if item.direction == SalesAdjustment.Direction.CASH_IN:
                    cash_in = amount
                else:
                    cash_out = amount
            else:
                user = item.user
                reference = f"AUDIT-{item.pk}"
                description = f"{item.get_action_display()} · {item.entity_repr}: {item.remarks}".strip(" ·")
                status = item.entity_type
            add_row(
                kind, event_date, item.created_at, reference, description,
                branch.name if hasattr(branch, "name") else branch,
                user, amount, cash_in, cash_out, status,
            )

    rows.sort(key=lambda row: (
        row["date"], row["time"], row["type"], row["reference"],
    ), reverse=True)
    rows = rows[start:end]

    branch_analysis = {}
    branch_fields = {
        "payment": ("loan__branch", "amount_received"),
        "refund": ("client__branch", "amount"),
        "credit": ("client__branch", "amount"),
        "injection": ("branch", "amount"),
        "expense": ("branch", "amount"),
        "loan": ("branch", "principal_amount"),
        "client": ("branch", None),
        "bank": ("branch", "amount"),
        "adjustment": ("branch", "amount"),
        "audit": (None, None),
    }
    for kind, (queryset, _date_field) in selected_querysets.items():
        field, amount_field = branch_fields.get(kind, (None, None))
        if not field:
            continue
        branch_id_key = f"{field}__pk"
        branch_name_key = f"{field}__name"
        aggregates = {"count": Count("pk")}
        if amount_field:
            aggregates["amount"] = Sum(Abs(amount_field) if kind == "refund" else amount_field)
        summary = queryset.values(branch_id_key, branch_name_key).annotate(**aggregates)
        for item in summary:
            key = item[branch_id_key]
            row = branch_analysis.setdefault(key, {
                "name": item[branch_name_key] or "Unassigned",
                "count": 0,
                "amount": ZERO,
            })
            row["count"] += item["count"]
            if amount_field:
                row["amount"] += item["amount"] or ZERO
    branch_analysis = sorted(
        branch_analysis.values(), key=lambda row: (row["count"], row["name"]), reverse=True,
    )

    performance_rows = _staff_rows(request, date_from, date_to)
    query_params = request.GET.copy()
    query_params.pop("page", None)
    filter_query = query_params.urlencode()
    branch_options = Branch.objects.filter(is_active=True).order_by("name")
    if not request.user.can("can_view_all_branches"):
        branch_options = branch_options.filter(pk=request.user.branch_id)
    return render(request, "reports/company_ledger.html", {
        "rows": rows,
        "date_from": date_from,
        "date_to": date_to,
        "today": today,
        "category": category,
        "categories": visible_categories,
        "category_label": visible_categories.get(category, "All record types"),
        "search": search,
        "total_count": total_count,
        "page_number": page_number,
        "page_count": page_count,
        "has_previous": page_number > 1,
        "has_next": page_number < page_count,
        "previous_page": page_number - 1,
        "next_page": page_number + 1,
        "filter_query": filter_query,
        "branch_id": "" if branch_id is None else str(branch_id),
        "branch_options": branch_options,
        "branch_analysis": branch_analysis,
        "performance_rows": performance_rows,
        "can_record_adjustment": request.user.can("can_manage_expenses"),
        "form_errors": errors,
        "form_values": form_values,
    })


@_require_manager
def report_index(request):
    return render(request, "reports/index.html")


# ---------------------------------------------------------------------------
# Loan book
# ---------------------------------------------------------------------------
@_require_manager
def loan_book(request):
    loans = list(_loan_book_qs(request))
    return render(request, "reports/loan_book.html", {
        "loans": loans,
        "total_principal": sum((l.principal_amount for l in loans), ZERO),
        "total_outstanding": sum((l.outstanding_balance for l in loans), ZERO),
        "total_paid": sum((l.total_paid for l in loans), ZERO),
        "products": LoanProduct.objects.filter(is_active=True),
        "date_from": request.GET.get("date_from", ""),
        "date_to": request.GET.get("date_to", ""),
        "product_id": request.GET.get("product", ""),
    })


@_require_manager
def loan_book_download(request):
    loans = list(_loan_book_qs(request))
    date_from = request.GET.get("date_from", "")
    date_to = request.GET.get("date_to", "")
    period = f"{date_from or 'Start'} to {date_to or 'End'}" if (date_from or date_to) else "All Loans"

    body_rows = [
        [p(l.loan_number), p(l.client.full_name), p(l.product.name), p(_ugx(l.principal_amount)),
         p(_ugx(l.outstanding_balance)), p(_ugx(l.total_paid)), p(l.get_status_display())]
        for l in loans
    ]
    totals_row = [
        p(""), p(""), p("<b>TOTAL</b>", CELL_BOLD),
        p(f"<b>{_ugx(sum((l.principal_amount for l in loans), ZERO))}</b>", CELL_BOLD),
        p(f"<b>{_ugx(sum((l.outstanding_balance for l in loans), ZERO))}</b>", CELL_BOLD),
        p(f"<b>{_ugx(sum((l.total_paid for l in loans), ZERO))}</b>", CELL_BOLD), p(""),
    ]
    return build_report_pdf(
        request, filename=f"LoanBook-{_today()}.pdf", title="Loan Book Report",
        subtitle=f"Period: {period} | {len(loans)} loans | Generated: {_stamp()}",
        landscape=True,
        sections=[{
            "heading": None,
            "head_row": ["Loan #", "Client", "Product", "Principal", "Outstanding", "Paid", "Status"],
            "col_widths": [28*mm, 60*mm, 40*mm, 34*mm, 34*mm, 34*mm, 30*mm],
            "body_rows": body_rows, "totals_row": totals_row,
        }],
    )


# ---------------------------------------------------------------------------
# Collections by cashier
# ---------------------------------------------------------------------------
@_require_manager
def collections_report(request):
    return render(request, "reports/collections.html", _collections_data(request))


@_require_manager
def collections_download(request):
    d = _collections_data(request)
    body_rows = [
        [p(r["user"].get_full_name() if r["user"] else "Unknown"), p(r["count"]), p(_ugx(r["total"]))]
        for r in d["cashier_totals"]
    ]
    totals_row = [p("<b>TOTAL</b>", CELL_BOLD), p(f"<b>{d['payment_count']}</b>", CELL_BOLD),
                  p(f"<b>{_ugx(d['grand_total'])}</b>", CELL_BOLD)]
    return build_report_pdf(
        request, filename=f"Collections-{_today()}.pdf", title="Collections by Cashier",
        subtitle=f"Period: {d['date_from']} to {d['date_to']} | Generated: {_stamp()}",
        sections=[{"heading": None, "head_row": ["Cashier", "Payments", "Total Collected"],
                   "col_widths": [90*mm, 40*mm, 50*mm], "body_rows": body_rows, "totals_row": totals_row}],
    )


# ---------------------------------------------------------------------------
# Overdue installments
# ---------------------------------------------------------------------------
@_require_manager
def overdue_report(request):
    return render(request, "reports/overdue.html", _overdue_data(request))


@_require_manager
def overdue_download(request):
    d = _overdue_data(request)
    body_rows = [
        [p(r["entry"].loan.loan_number), p(r["entry"].loan.client.full_name),
         p(r["entry"].due_date.isoformat()), p(r["days_overdue"]), p(_ugx(r["overdue_amount"]))]
        for r in d["rows"]
    ]
    totals_row = [p(""), p(""), p(""), p("<b>TOTAL</b>"), p(f"<b>{_ugx(d['total_overdue'])}</b>")]
    return build_report_pdf(
        request, filename=f"Overdue-{d['today']}.pdf", title="Overdue Installments Report",
        subtitle=f"As at {d['today']:%d %b %Y} | Generated: {_stamp()}",
        sections=[{"heading": None,
                   "head_row": ["Loan #", "Client", "Due Date", "Days Overdue", "Amount Overdue"],
                   "col_widths": [28*mm, 55*mm, 28*mm, 28*mm, 41*mm],
                   "body_rows": body_rows, "totals_row": totals_row}],
    )


# ---------------------------------------------------------------------------
# Monthly income statement
# ---------------------------------------------------------------------------
@_require_manager
def income_statement(request):
    return render(request, "reports/income_statement.html", _income_data(request))


@_require_manager
def income_download(request):
    d = _income_data(request)
    body_rows = [
        [p("Interest Income"), p(_ugx(d["total_interest"]))],
        [p("Penalty Income"), p(_ugx(d["total_penalties"]))],
        [p("Processing Fee Income"), p(_ugx(d["total_fees"]))],
        [p("<b>Total Income</b>", CELL_BOLD), p(f"<b>{_ugx(d['total_income'])}</b>", CELL_BOLD)],
        [p("Operating Expenses (approved)"), p(f"- {_ugx(d['total_expenses'])}")],
        [p("<b>Net Income</b>", CELL_BOLD), p(f"<b>{_ugx(d['net_income'])}</b>", CELL_BOLD)],
        [p("Principal Recovered"), p(_ugx(d["total_principal"]))],
        [p("Cash Received (excl. credit transfers)"), p(_ugx(d["total_received"]))],
        [p("Client credit applied / (held) — reconciles cash to allocations"), p(_ugx(d["credit_adjustment"]))],
        [p("Loans Disbursed (count)"), p(d["loan_count"])],
        [p("Total Principal Disbursed"), p(_ugx(d["total_disbursed"]))],
    ]
    return build_report_pdf(
        request, filename=f"IncomeStatement-{d['month_str']}.pdf", title="Monthly Income Statement",
        subtitle=f"Period: {d['period_start']:%d %b %Y} – {d['period_end']:%d %b %Y} | Generated: {_stamp()}",
        sections=[{"heading": None, "head_row": ["Item", "Amount"],
                   "col_widths": [110*mm, 60*mm], "body_rows": body_rows}],
    )


# ---------------------------------------------------------------------------
# Cash in / cash out
# ---------------------------------------------------------------------------
@_require_manager
def cash_flow_report(request):
    return render(request, "reports/cash_flow.html", _cash_flow_data(request))


@_require_manager
def cash_flow_download(request):
    d = _cash_flow_data(request)
    body_rows = [
        [p(r["date"].isoformat()), p(timezone.localtime(r["time"]).strftime("%H:%M:%S")), p(r["type"]),
         p(r["client"].full_name if r["client"] else "—"), p(r["description"]),
         p(_ugx(r["cash_in"])), p(_ugx(r["cash_out"]))]
        for r in d["ledger"]
    ]
    totals_row = [p(""), p(""), p(""), p(""), p("<b>TOTAL</b>"),
                  p(f"<b>{_ugx(d['total_cash_in'])}</b>"), p(f"<b>{_ugx(d['total_cash_out'])}</b>")]
    return build_report_pdf(
        request, filename=f"CashFlow-{_today()}.pdf", title="Cash In / Cash Out Ledger",
        subtitle=f"Period: {d['date_from']} to {d['date_to']} | Generated: {_stamp()}",
        landscape=True,
        sections=[{"heading": None,
                   "head_row": ["Date", "Time", "Type", "Client", "Description", "Cash In", "Cash Out"],
                   "col_widths": [24*mm, 18*mm, 30*mm, 55*mm, 70*mm, 35*mm, 35*mm],
                   "body_rows": body_rows, "totals_row": totals_row}],
    )


# ---------------------------------------------------------------------------
# Capital and fund movement
# ---------------------------------------------------------------------------
@_require_manager
def fund_movement_report(request):
    return render(request, "reports/fund_movement.html", _fund_movement_data(request))


@_require_manager
def fund_movement_pdf(request):
    data = _fund_movement_data(request)
    summary_rows = [
        [p("Opening balance"), p(_ugx(data["opening_balance"]))],
        [p("Capital injections"), p(_ugx(data["total_injections"]))],
        [p("Client repayments"), p(_ugx(data["total_repayments"]))],
        [p("Payment reversals"), p(_ugx(data["total_reversals"]))],
        [p("Client refunds"), p(_ugx(data["total_refunds"]))],
        [p("Processing fees"), p(_ugx(data["total_processing_fees"]))],
        [p("Sales adjustments — cash in"), p(_ugx(data["total_adjustment_in"]))],
        [p("Sales adjustments — cash out"), p(_ugx(data["total_adjustment_out"]))],
        [p("Total inflows"), p(_ugx(data["total_inflows"]))],
        [p("Loan principal disbursed"), p(_ugx(data["total_disbursed"]))],
        [p("Approved expenses"), p(_ugx(data["total_expenses"]))],
        [p("Total outflows"), p(_ugx(data["total_outflows"]))],
        [p("Net movement"), p(_ugx(data["net_movement"]))],
        [p("Closing balance"), p(_ugx(data["closing_balance"]))],
        [p("Capital injection count"), p(data["injection_count"])],
        [p("Repayment count"), p(data["payment_count"])],
        [p("Loan disbursement count"), p(data["disbursement_count"])],
        [p("Approved expense count"), p(data["expense_count"])],
        [p("Sales adjustment count"), p(data["adjustment_count"])],
        [p("Payment reversal count"), p(data["reversal_count"])],
        [p("Client refund count"), p(data["refund_count"])],
    ]
    body_rows = [
        [
            p(row["date"].isoformat()),
            p(timezone.localtime(row["time"]).strftime("%H:%M:%S")),
            p(escape(row["type"])),
            p(escape(row["counterparty"])),
            p(escape(row["method"])),
            p(escape(row["reference"])),
            p(escape(row["description"])),
            p(_ugx(row["cash_in"])),
            p(_ugx(row["cash_out"])),
            p(_ugx(row["balance_after"])),
        ]
        for row in data["ledger"]
    ]
    return build_report_pdf(
        request,
        filename=f"Fund-Movement-{data['date_from']}-{data['date_to']}.pdf",
        title="Capital & Fund Movement Report",
        subtitle=(
            f"Period: {data['date_from']} to {data['date_to']} | "
            f"Branch: {data['branch_label']} | Generated: {_stamp()}"
        ),
        landscape=True,
        sections=[
            {"heading": "Fund Reconciliation", "head_row": ["Summary", "Amount"],
             "col_widths": [110*mm, 70*mm], "body_rows": summary_rows},
            {"heading": "Transaction Ledger",
             "head_row": ["Date", "Time", "Type", "Counterparty", "Channel", "Reference", "Description", "In", "Out", "Balance"],
             "col_widths": [20*mm, 14*mm, 24*mm, 30*mm, 21*mm, 26*mm, 55*mm, 28*mm, 28*mm, 27*mm],
             "body_rows": body_rows},
        ],
    )


def _excel_safe_text(value):
    text = str(value or "")
    return f"'{text}" if text.startswith(("=", "+", "-", "@")) else text


@_require_manager
def fund_movement_excel(request):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    data = _fund_movement_data(request)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Fund Summary"
    summary.append(["Capital & Fund Movement Report"])
    summary.append(["Period", f"{data['date_from']} to {data['date_to']}"])
    summary.append(["Branch", data["branch_label"]])
    summary.append(["Generated", _stamp()])
    summary.append([])
    summary.append(["Measure", "Amount (UGX)"])
    for label, key in (
        ("Opening balance", "opening_balance"),
        ("Capital injections", "total_injections"),
        ("Client repayments", "total_repayments"),
        ("Payment reversals", "total_reversals"),
        ("Client refunds", "total_refunds"),
        ("Processing fees", "total_processing_fees"),
        ("Sales adjustments — cash in", "total_adjustment_in"),
        ("Sales adjustments — cash out", "total_adjustment_out"),
        ("Total inflows", "total_inflows"),
        ("Loan principal disbursed", "total_disbursed"),
        ("Approved expenses", "total_expenses"),
        ("Total outflows", "total_outflows"),
        ("Net movement", "net_movement"),
        ("Closing balance", "closing_balance"),
    ):
        summary.append([label, float(data[key])])
    summary.append(["Capital injection count", data["injection_count"]])
    summary.append(["Repayment count", data["payment_count"]])
    summary.append(["Loan disbursement count", data["disbursement_count"]])
    summary.append(["Approved expense count", data["expense_count"]])
    summary.append(["Sales adjustment count", data["adjustment_count"]])
    summary.append(["Payment reversal count", data["reversal_count"]])
    summary.append(["Client refund count", data["refund_count"]])

    ledger_sheet = workbook.create_sheet("Movement Ledger")
    ledger_sheet.append([
        "Date", "Time", "Type", "Counterparty", "Payment Method", "Reference",
        "Description", "Branch", "Funds In (UGX)", "Funds Out (UGX)", "Balance After (UGX)",
    ])
    for row in data["ledger"]:
        ledger_sheet.append([
            row["date"],
            timezone.localtime(row["time"]).replace(tzinfo=None),
            _excel_safe_text(row["type"]),
            _excel_safe_text(row["counterparty"]),
            _excel_safe_text(row["method"]),
            _excel_safe_text(row["reference"]),
            _excel_safe_text(row["description"]),
            _excel_safe_text(row["branch"]),
            float(row["cash_in"]),
            float(row["cash_out"]),
            float(row["balance_after"]),
        ])

    header_fill = PatternFill("solid", fgColor="163B50")
    for sheet in (summary, ledger_sheet):
        sheet.freeze_panes = "A2" if sheet is ledger_sheet else "A7"
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
        for column_cells in sheet.columns:
            column = column_cells[0].column
            longest = max((len(str(cell.value or "")) for cell in column_cells), default=0)
            sheet.column_dimensions[get_column_letter(column)].width = min(max(longest + 2, 12), 48)
    summary[6][0].font = Font(bold=True)
    summary[6][1].font = Font(bold=True)
    for row in range(7, 21):
        summary.cell(row, 2).number_format = '#,##0.00;[Red]-#,##0.00'
    for row in range(2, ledger_sheet.max_row + 1):
        ledger_sheet.cell(row, 1).number_format = "dd mmm yyyy"
        ledger_sheet.cell(row, 2).number_format = "hh:mm:ss"
        for column in range(9, 12):
            ledger_sheet.cell(row, column).number_format = '#,##0.00;[Red]-#,##0.00'
    ledger_sheet.auto_filter.ref = ledger_sheet.dimensions
    for cell in ledger_sheet[1]:
        cell.alignment = Alignment(wrap_text=True)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        f"attachment; filename=Fund-Movement-{data['date_from']}-{data['date_to']}.xlsx"
    )
    return response


# ---------------------------------------------------------------------------
# Disbursements
# ---------------------------------------------------------------------------
def _disbursements_data(request):
    date_from, date_to = _period(request)
    product_id = request.GET.get("product", "")
    loans = list(_disbursed_qs(request, date_from, date_to, product_id)
                 .select_related("client", "product", "applied_by", "reviewed_by")
                 .order_by("-disbursement_date"))
    return {
        "loans": loans,
        "total_disbursed": sum((l.principal_amount for l in loans), ZERO),
        "total_interest_exp": sum((l.total_interest for l in loans), ZERO),
        "total_fees": _fees(loans),
        "loan_count": len(loans),
        "product_id": product_id, "date_from": date_from, "date_to": date_to,
    }


@_require_manager
def disbursements_report(request):
    d = _disbursements_data(request)
    d["products"] = LoanProduct.objects.filter(is_active=True)
    return render(request, "reports/disbursements.html", d)


@_require_manager
def disbursements_download(request):
    d = _disbursements_data(request)
    body_rows = [
        [p(l.loan_number), p(l.client.full_name), p(l.product.name),
         p(l.applied_by.get_full_name() if l.applied_by else "—"),
         p(_ugx(l.principal_amount)), p(_ugx(l.effective_processing_fee)),
         p(l.disbursement_date.isoformat() if l.disbursement_date else "—")]
        for l in d["loans"]
    ]
    totals_row = [p(""), p(""), p(""), p("<b>TOTAL</b>"), p(f"<b>{_ugx(d['total_disbursed'])}</b>"),
                  p(f"<b>{_ugx(d['total_fees'])}</b>"), p("")]
    return build_report_pdf(
        request, filename=f"Disbursements-{_today()}.pdf", title="Disbursements Report",
        subtitle=f"Period: {d['date_from']} to {d['date_to']} | Total Interest Expected: "
                 f"{_ugx(d['total_interest_exp'])} | Generated: {_stamp()}",
        landscape=True,
        sections=[{"heading": None,
                   "head_row": ["Loan #", "Client", "Product", "Loan Officer", "Principal", "Fees", "Disbursed"],
                   "col_widths": [24*mm, 50*mm, 35*mm, 40*mm, 32*mm, 28*mm, 28*mm],
                   "body_rows": body_rows, "totals_row": totals_row}],
    )


# ---------------------------------------------------------------------------
# Repayments
# ---------------------------------------------------------------------------
def _repayments_data(request):
    date_from, date_to = _period(request)
    method = request.GET.get("method", "")
    qs = (_payments_qs(request, date_from, date_to, method)
          .select_related("loan__client", "loan__product", "recorded_by").order_by("-payment_date"))
    return {
        "payments": qs,
        "total_received": _sum(qs.cash_receipts(), "amount_received"),
        "total_principal": _sum(qs, "principal_paid"),
        "total_interest": _sum(qs, "interest_paid"),
        "total_penalty": _sum(qs, "penalty_paid"),
        "payment_count": qs.count(),
        "date_from": date_from, "date_to": date_to, "method_filter": method,
    }


@_require_manager
def repayments_report(request):
    d = _repayments_data(request)
    d["methods"] = Payment.PaymentMethod.choices
    return render(request, "reports/repayments.html", d)


@_require_manager
def repayments_download(request):
    d = _repayments_data(request)
    body_rows = [
        [p(pm.payment_date.isoformat()), p(pm.loan.loan_number), p(pm.loan.client.full_name),
         p(pm.get_payment_method_display()), p(_ugx(pm.amount_received)),
         p(_ugx(pm.principal_paid)), p(_ugx(pm.interest_paid)), p(_ugx(pm.penalty_paid))]
        for pm in d["payments"]
    ]
    totals_row = [p(""), p(""), p(""), p("<b>TOTAL</b>"), p(f"<b>{_ugx(d['total_received'])}</b>"),
                  p(f"<b>{_ugx(d['total_principal'])}</b>"), p(f"<b>{_ugx(d['total_interest'])}</b>"),
                  p(f"<b>{_ugx(d['total_penalty'])}</b>")]
    return build_report_pdf(
        request, filename=f"Repayments-{_today()}.pdf", title="Repayments Report",
        subtitle=f"Period: {d['date_from']} to {d['date_to']} | Cash total excludes internal credit "
                 f"transfers | Generated: {_stamp()}",
        landscape=True,
        sections=[{"heading": None,
                   "head_row": ["Date", "Loan #", "Client", "Method", "Received", "Principal", "Interest", "Penalty"],
                   "col_widths": [22*mm, 22*mm, 45*mm, 25*mm, 28*mm, 28*mm, 25*mm, 25*mm],
                   "body_rows": body_rows, "totals_row": totals_row}],
    )


# ---------------------------------------------------------------------------
# Defaulted / written-off
# ---------------------------------------------------------------------------
@_require_manager
def defaulted_loans_report(request):
    return render(request, "reports/defaulted.html", _defaulted_data(request))


@_require_manager
def defaulted_download(request):
    d = _defaulted_data(request)
    defaulted_rows = [
        [p(l.loan_number), p(l.client.full_name), p(l.product.name), p(l.get_status_display()),
         p(_ugx(l.principal_amount)),
         p(_ugx(l.written_off_amount if l.status == "WRITTEN_OFF" else l.outstanding_balance))]
        for l in d["loans"]
    ]
    defaulted_totals = [p(""), p(""), p(""), p("<b>TOTAL</b>"), p(f"<b>{_ugx(d['total_defaulted'])}</b>"),
                        p(f"<b>{_ugx(d['total_outstanding'] + d['total_written_off'])}</b>")]
    par90_rows = [[p(l.loan_number), p(l.client.full_name), p(l.product.name), p(_ugx(l.outstanding_balance))]
                  for l in d["par90_loans"]]
    return build_report_pdf(
        request, filename=f"Defaulted-{d['today']}.pdf", title="Defaulted & Written-Off Loans",
        subtitle=f"As at {d['today']:%d %b %Y} | Generated: {_stamp()}", landscape=True,
        sections=[
            {"heading": "Defaulted / Written Off Loans (amount = outstanding, or amount written off)",
             "head_row": ["Loan #", "Client", "Product", "Status", "Principal", "Outstanding / Written off"],
             "col_widths": [24*mm, 55*mm, 35*mm, 30*mm, 35*mm, 45*mm],
             "body_rows": defaulted_rows, "totals_row": defaulted_totals},
            {"heading": "Active Loans Over 90 Days Overdue (PAR90)",
             "head_row": ["Loan #", "Client", "Product", "Outstanding"],
             "col_widths": [24*mm, 70*mm, 45*mm, 35*mm], "body_rows": par90_rows},
        ],
    )


# ---------------------------------------------------------------------------
# Closed loans
# ---------------------------------------------------------------------------
def _closed_data(request):
    qs = Loan.objects.filter(status="COMPLETED").select_related("client", "product").order_by("-completion_date")
    date_from = request.GET.get("date_from", "")
    date_to = request.GET.get("date_to", "")
    product_id = request.GET.get("product", "")
    if date_from:
        qs = qs.filter(completion_date__gte=date_from)
    if date_to:
        qs = qs.filter(completion_date__lte=date_to)
    if product_id:
        qs = qs.filter(product_id=product_id)
    qs = _branch(qs, request)
    return {
        "loans": qs,
        "total_principal": _sum(qs, "principal_amount"),
        "total_interest": _sum(qs, "total_interest"),
        "total_collected": _sum(qs, "total_paid"),
        "loan_count": qs.count(),
        "product_id": product_id, "date_from": date_from, "date_to": date_to,
    }


@_require_manager
def closed_loans_report(request):
    d = _closed_data(request)
    d["products"] = LoanProduct.objects.filter(is_active=True)
    return render(request, "reports/closed_loans.html", d)


@_require_manager
def closed_loans_download(request):
    d = _closed_data(request)
    body_rows = [
        [p(l.loan_number), p(l.client.full_name), p(l.product.name),
         p(l.completion_date.isoformat() if l.completion_date else "—"),
         p(_ugx(l.principal_amount)), p(_ugx(l.total_interest)), p(_ugx(l.total_paid))]
        for l in d["loans"]
    ]
    totals_row = [p(""), p(""), p(""), p("<b>TOTAL</b>"), p(f"<b>{_ugx(d['total_principal'])}</b>"),
                  p(f"<b>{_ugx(d['total_interest'])}</b>"), p(f"<b>{_ugx(d['total_collected'])}</b>")]
    return build_report_pdf(
        request, filename=f"ClosedLoans-{_today()}.pdf", title="Closed (Fully Repaid) Loans",
        subtitle=f"Period: {d['date_from'] or 'Any'} to {d['date_to'] or 'Any'} | Generated: {_stamp()}",
        landscape=True,
        sections=[{"heading": None,
                   "head_row": ["Loan #", "Client", "Product", "Completed", "Principal", "Interest Earned", "Total Collected"],
                   "col_widths": [22*mm, 45*mm, 30*mm, 25*mm, 30*mm, 33*mm, 33*mm],
                   "body_rows": body_rows, "totals_row": totals_row}],
    )


# ---------------------------------------------------------------------------
# Portfolio at risk
# ---------------------------------------------------------------------------
@_require_manager
def par_report(request):
    d = _par_data(request)
    ctx = {"active_portfolio": d["portfolio"], "today": d["today"], "par_data": []}
    for b in d["buckets"]:
        ctx[b["key"]] = b["rows"]
        ctx[f"{b['key']}_total"] = b["at_risk"]      # outstanding balance at risk
        ctx[f"{b['key']}_pct"] = b["pct"]
        ctx["par_data"].append((b["rows"], b["rows"], b["label"], b["css"], b["at_risk"], b["pct"]))
    return render(request, "reports/par.html", ctx)


@_require_manager
def par_download(request):
    d = _par_data(request)
    sections = []
    for b in d["buckets"]:
        sections.append({
            "heading": f"{b['label']}  —  {_ugx(b['at_risk'])} outstanding at risk  ({b['pct']}% of portfolio)",
            "head_row": ["Loan #", "Client", "Product", "Outstanding", "Overdue Amount", "Days Overdue"],
            "col_widths": [24*mm, 55*mm, 35*mm, 35*mm, 35*mm, 25*mm],
            "body_rows": [
                [p(r["loan"].loan_number), p(r["loan"].client.full_name), p(r["loan"].product.name),
                 p(_ugx(r["loan"].outstanding_balance)), p(_ugx(r["overdue_amount"])), p(r["days_overdue"])]
                for r in b["rows"]
            ],
        })
    return build_report_pdf(
        request, filename=f"PAR-{d['today']}.pdf", title="Portfolio at Risk (PAR)",
        subtitle=f"As at {d['today']:%d %b %Y} | Active Portfolio: {_ugx(d['portfolio'])} | Generated: {_stamp()}",
        landscape=True, sections=sections,
    )


# ---------------------------------------------------------------------------
# Client statement
# ---------------------------------------------------------------------------
@_require_manager
def client_statement(request):
    """Per-client loan & payment statement."""
    client_id = request.GET.get("client", "")
    client = None
    loans = []
    payments = []

    all_clients = scope_to_branch(Client.objects.filter(is_active=True), request.user).order_by("last_name", "first_name")

    if client_id:
        client = get_object_or_404(Client, pk=client_id)
        if not can_access_branch_object(request.user, client):
            messages.error(request, "That client belongs to a different branch.")
            return redirect("reports:client_statement")
        loans = Loan.objects.filter(client=client).select_related("product").order_by("-application_date")
        payments = Payment.objects.filter(client=client, status="ALLOCATED").select_related("loan").order_by("-payment_date")

    return render(request, "reports/client_statement.html", {
        "all_clients": all_clients, "client": client, "loans": loans,
        "payments": payments, "client_id": client_id,
    })


@_require_manager
def client_statement_download(request):
    client_id = request.GET.get("client", "")
    client = get_object_or_404(Client, pk=client_id) if client_id else None
    today = _today()

    if client and not can_access_branch_object(request.user, client):
        messages.error(request, "That client belongs to a different branch.")
        return redirect("reports:client_statement")

    if not client:
        return build_report_pdf(
            request, filename=f"ClientStatement-{today}.pdf", title="Client Account Statement",
            subtitle="No client selected.",
            sections=[{"heading": None, "head_row": ["—"], "col_widths": [170*mm], "body_rows": []}],
        )

    loans = Loan.objects.filter(client=client).select_related("product").order_by("-application_date")
    payments = Payment.objects.filter(client=client, status="ALLOCATED").select_related("loan").order_by("-payment_date")

    loan_rows = [[p(l.loan_number), p(l.product.name), p(_ugx(l.principal_amount)),
                  p(_ugx(l.outstanding_balance)), p(l.get_status_display())] for l in loans]
    payment_rows = [[p(pm.payment_date.isoformat()), p(pm.loan.loan_number), p(_ugx(pm.amount_received)),
                     p(pm.get_payment_method_display())] for pm in payments]

    return build_report_pdf(
        request, filename=f"ClientStatement-{client.full_name.replace(' ', '')}-{today}.pdf",
        title=f"Client Account Statement — {client.full_name}",
        subtitle=f"Phone: {client.phone_primary} | Generated: {_stamp()}",
        sections=[
            {"heading": "Loans", "head_row": ["Loan #", "Product", "Principal", "Outstanding", "Status"],
             "col_widths": [26*mm, 40*mm, 35*mm, 35*mm, 34*mm], "body_rows": loan_rows},
            {"heading": "Payment History", "head_row": ["Date", "Loan #", "Amount", "Method"],
             "col_widths": [35*mm, 35*mm, 50*mm, 50*mm], "body_rows": payment_rows},
        ],
    )


# ---------------------------------------------------------------------------
# Staff performance
# ---------------------------------------------------------------------------
@_require_manager
def staff_performance_report(request):
    date_from, date_to = _period(request)
    branch_id = _effective_branch_id(request)
    rows = _staff_rows(request, date_from, date_to)
    User = get_user_model()
    return render(request, "reports/staff_performance.html", {
        "rows": rows, "date_from": date_from, "date_to": date_to, "today": _today(),
        "branches": Branch.objects.filter(is_active=True).order_by("name"),
        "selected_branch": str(branch_id) if branch_id and branch_id != -1 else "",
        "roles": User.Role.choices, "selected_role": request.GET.get("role", ""),
        "staff_members": User.objects.filter(is_active=True).order_by("last_name", "first_name"),
        "selected_staff": request.GET.get("staff", ""),
    })


@_require_manager
def staff_performance_download(request):
    date_from, date_to = _period(request)
    rows = _staff_rows(request, date_from, date_to)
    body_rows = [
        [p(r["user"].get_full_name()), p(r["disbursed_count"]), p(_ugx(r["disbursed_amount"])),
         p(_ugx(r["collected_amount"])), p(r["overdue_count"]), p(r["defaulted_count"]),
         p(f"{r['quality_score']}%")]
        for r in rows
    ]
    return build_report_pdf(
        request, filename=f"StaffPerformance-{_today()}.pdf", title="Staff Performance Report",
        subtitle=f"Period: {date_from} to {date_to} | Generated: {_stamp()}", landscape=True,
        sections=[{"heading": None,
                   "head_row": ["Staff", "Loans Disbursed", "Amount Disbursed", "Amount Collected",
                                "Overdue Loans", "Defaulted", "Quality Score"],
                   "col_widths": [40*mm, 25*mm, 35*mm, 35*mm, 25*mm, 22*mm, 25*mm],
                   "body_rows": body_rows}],
    )
