from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from django.core.exceptions import ValidationError
from django.test import RequestFactory, SimpleTestCase
from django.utils import timezone
from openpyxl import load_workbook

from accounts.models import SalesAdjustment
from . import views


def _sample_report_data():
    return {
        "date_from": date(2026, 1, 1),
        "date_to": date(2026, 3, 31),
        "branch_label": "All branches",
        "opening_balance": Decimal("100000"),
        "total_injections": Decimal("500000"),
        "total_repayments": Decimal("250000"),
        "total_reversals": Decimal("0"),
        "total_refunds": Decimal("0"),
        "total_processing_fees": Decimal("10000"),
        "total_adjustment_in": Decimal("0"),
        "total_adjustment_out": Decimal("0"),
        "total_inflows": Decimal("760000"),
        "total_disbursed": Decimal("400000"),
        "total_expenses": Decimal("60000"),
        "total_outflows": Decimal("460000"),
        "net_movement": Decimal("300000"),
        "closing_balance": Decimal("400000"),
        "payment_count": 1,
        "injection_count": 1,
        "expense_count": 1,
        "adjustment_count": 0,
        "reversal_count": 0,
        "refund_count": 0,
        "disbursement_count": 1,
        "ledger": [{
            "date": date(2026, 1, 2),
            "time": timezone.make_aware(datetime(2026, 1, 2, 9, 30)),
            "type": "Capital injection",
            "counterparty": "Investor",
            "method": "Bank Transfer",
            "reference": "INJ-1",
            "description": "Seed funding",
            "branch": "HQ",
            "cash_in": Decimal("500000"),
            "cash_out": Decimal("0"),
            "balance_after": Decimal("600000"),
        }],
    }


class FundMovementExportTests(SimpleTestCase):
    def setUp(self):
        self.request = RequestFactory().get("/reports/fund-movement/")
        self.request.session = {}
        self.request.user = SimpleNamespace(
            is_authenticated=True,
            can=lambda _permission: True,
        )

    @patch("reports.views._fund_movement_data", return_value=_sample_report_data())
    def test_pdf_export_returns_a_pdf(self, _report_data):
        response = views.fund_movement_pdf(self.request)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))

    @patch("reports.views._fund_movement_data", return_value=_sample_report_data())
    def test_excel_export_contains_summary_and_ledger(self, _report_data):
        response = views.fund_movement_excel(self.request)
        workbook = load_workbook(BytesIO(response.content), data_only=True)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            workbook.sheetnames,
            ["Fund Summary", "Movement Ledger"],
        )
        self.assertEqual(workbook["Fund Summary"]["B20"].value, 400000)
        self.assertEqual(workbook["Movement Ledger"]["F2"].value, "INJ-1")


class SalesAdjustmentValidationTests(SimpleTestCase):
    def test_amount_must_be_positive(self):
        adjustment = SalesAdjustment(
            reference="CORR-1",
            direction=SalesAdjustment.Direction.CASH_IN,
            amount=Decimal("0"),
        )

        with self.assertRaises(ValidationError):
            adjustment.clean_fields(exclude=("reference", "branch", "created_by"))

    def test_valid_direction_and_positive_amount_pass_field_validation(self):
        adjustment = SalesAdjustment(
            reference="CORR-2",
            direction=SalesAdjustment.Direction.CASH_OUT,
            amount=Decimal("25.00"),
        )

        adjustment.clean_fields(exclude=("reference", "branch", "created_by"))